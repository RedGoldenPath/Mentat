"""Text extraction for user attachments (drag-and-drop into the composer).

Borrowed from v1 (`brain_core.extract_text`): txt/md/rst/html, pdf (pymupdf text layer,
Apple-Vision OCR fallback for scans), epub (ebooklib), docx (python-docx). CSV/JSON/code are
read as plain text. This mirrors what Qwen/DeepSeek/Kimi/ChatGPT do server-side: parse the
uploaded file to text so the model can read it.
"""
from __future__ import annotations
import html
import json
import re
from pathlib import Path

MAX_CHARS = 400_000          # hard cap on extracted text handed to the engine
_TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml",
             ".py", ".js", ".ts", ".swift", ".c", ".cpp", ".h", ".java", ".go", ".rs",
             ".sh", ".sql", ".html", ".htm", ".xml", ".log"}


def _extract_txt(path):
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return Path(path).read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return Path(path).read_text(errors="ignore")


def _ocr_page_text(page, dpi=200):
    """Apple Vision OCR of one PDF page (rendered to an image). macOS only."""
    import Vision
    pix = page.get_pixmap(dpi=dpi)
    request = Vision.VNRecognizeTextRequest.alloc().init()
    request.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
    request.setUsesLanguageCorrection_(True)
    handler = Vision.VNImageRequestHandler.alloc().initWithData_options_(pix.tobytes("png"), None)
    ok, _err = handler.performRequests_error_([request], None)
    if not ok:
        return ""
    results = sorted(request.results() or [],
                     key=lambda r: (-r.boundingBox().origin.y, r.boundingBox().origin.x))
    return "\n".join(r.topCandidates_(1)[0].string() for r in results)


def _extract_pdf(path):
    import pymupdf
    doc = pymupdf.open(str(path))
    try:
        parts = []
        for index, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if not text:
                try:
                    text = _ocr_page_text(page)
                except Exception:
                    text = ""
            if text:
                parts.append(f"## Page {index}\n\n{text}")
        return "\n\n".join(parts)
    finally:
        doc.close()


def _strip_html(text):
    text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
    text = re.sub(r"</?(p|div|br|li|h[1-6]|blockquote|tr|ul|ol|table)[^>]*>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n")]
    out, blanks = [], 0
    for ln in lines:
        if ln:
            out.append(ln); blanks = 0
        elif out and blanks < 1:
            out.append(""); blanks += 1
    return "\n".join(out).strip()


def _extract_epub(path):
    import ebooklib
    from ebooklib import epub
    book = epub.read_epub(str(path))
    parts = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        getter = getattr(item, "get_body_content", None) or getattr(item, "get_content", None)
        if getter is None:
            continue
        try:
            content = getter() or b""
        except Exception:
            content = b""
        text = _strip_html(content.decode("utf-8", errors="ignore"))
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def _extract_docx(path):
    from docx import Document
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n\n".join(parts)


SUPPORTED = sorted(_TEXT_EXT | {".pdf", ".epub", ".docx"})


def extract_text(path, max_chars=MAX_CHARS):
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(str(path))
    suffix = p.suffix.lower()
    if suffix in _TEXT_EXT:
        text = _extract_txt(p)
    elif suffix == ".pdf":
        text = _extract_pdf(p)
    elif suffix == ".epub":
        text = _extract_epub(p)
    elif suffix == ".docx":
        text = _extract_docx(p)
    else:
        raise ValueError(f"unsupported file type: {suffix or '(none)'}")
    text = text or ""
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n[... truncated at {max_chars} characters ...]"
    return text
